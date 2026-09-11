from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = r"C:\DellApplications\NewAssign Prep\Sushma_Dhar_Enhanced_AI_DotNetLead_Resume.docx"

BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(37, 45, 55)
GRAY = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def bullet(doc, text):
    p = doc.add_paragraph(text, style="Resume Bullet")
    p.paragraph_format.keep_together = True
    return p


def role(doc, title, company, dates, location="Bengaluru, India"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{title} | {company}")
    r.bold = True
    r.font.color.rgb = DARK
    r.font.size = Pt(10.5)
    r = p.add_run(f"\n{dates} | {location}")
    r.italic = True
    r.font.color.rgb = GRAY
    r.font.size = Pt(9)


def section(doc, title):
    p = doc.add_paragraph(title, style="Resume Section")
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.45)
sec.bottom_margin = Inches(0.45)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.5)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.05

name_style = styles.add_style("Resume Name", WD_STYLE_TYPE.PARAGRAPH)
name_style.font.name = "Aptos Display"
name_style.font.size = Pt(22)
name_style.font.bold = True
name_style.font.color.rgb = BLUE
name_style.paragraph_format.space_after = Pt(1)
name_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

section_style = styles.add_style("Resume Section", WD_STYLE_TYPE.PARAGRAPH)
section_style.font.name = "Aptos Display"
section_style.font.size = Pt(11.5)
section_style.font.bold = True
section_style.font.color.rgb = BLUE
section_style.paragraph_format.space_before = Pt(9)
section_style.paragraph_format.space_after = Pt(3)
section_style.paragraph_format.keep_with_next = True
section_style.paragraph_format.border_bottom = True

bullet_style = styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
bullet_style.base_style = styles["List Bullet"]
bullet_style.font.name = "Aptos"
bullet_style.font.size = Pt(9.5)
bullet_style.font.color.rgb = DARK
bullet_style.paragraph_format.left_indent = Inches(0.2)
bullet_style.paragraph_format.first_line_indent = Inches(-0.15)
bullet_style.paragraph_format.space_after = Pt(2)
bullet_style.paragraph_format.line_spacing = 1.02

p = doc.add_paragraph("SUSHMA VAIDYA", style="Resume Name")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("PRINCIPAL SOFTWARE ENGINEER | TECHNICAL LEAD | .NET, JAVA, PYTHON & AI ENGINEERING")
r.bold = True
r.font.size = Pt(10.5)
r.font.color.rgb = DARK
p = doc.add_paragraph("Bengaluru, India  |  +91 9886843824  |  sush.dhar@gmail.com")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(6)
p.runs[0].font.color.rgb = GRAY
p.runs[0].font.size = Pt(9.5)

section(doc, "PROFESSIONAL SUMMARY")
p = doc.add_paragraph(
    "Principal Software Engineer and technical leader with 14+ years of experience designing, modernizing, "
    "and delivering enterprise applications across .NET, Java, Python, microservices, and data platforms. "
    "Proven success leading mission-critical order-processing and financial-validation systems at Dell "
    "Technologies, improving checkout performance by 40%, reducing latency by 35%, and lowering technical "
    "debt by 30%. Experienced in API architecture, event-driven systems, cloud-native delivery, security, "
    "CI/CD, performance engineering, stakeholder management, and mentoring. Currently expanding into AI "
    "engineering through hands-on work in advanced RAG, hybrid search, vector databases, reranking, SQL RAG, "
    "LLM APIs, and AI-assisted software development."
)
p.paragraph_format.space_after = Pt(4)

section(doc, "CORE COMPETENCIES")
table = doc.add_table(rows=0, cols=2)
table.autofit = True
skills = [
    ("Architecture & Leadership", "Solution architecture, technical leadership, design reviews, mentoring, stakeholder collaboration"),
    ("Backend Engineering", ".NET, ASP.NET, Java, Spring Boot, Python, REST APIs, microservices, event-driven architecture"),
    ("AI Engineering", "RAG, hybrid dense + BM25 search, Qdrant, cross-encoder reranking, SQL RAG, LLM APIs, prompt engineering"),
    ("Data & Integration", "SQL Server, MongoDB, SQLite, SSIS, SSRS, enterprise integrations, data warehousing"),
    ("Cloud & DevOps", "Cloud-native architecture, Kubernetes, CI/CD, PCF, OAuth, observability, performance optimization"),
    ("Engineering Practices", "Spec-driven development, secure coding, code reviews, refactoring, technical debt reduction, SIT/UAT"),
]
for label, value in skills:
    cells = table.add_row().cells
    cells[0].width = Inches(1.65)
    cells[1].width = Inches(5.7)
    cells[0].paragraphs[0].add_run(label).bold = True
    cells[1].paragraphs[0].add_run(value)
    for cell in cells:
        cell.vertical_alignment = 1
        cell.paragraphs[0].paragraph_format.space_after = Pt(1)
        cell.paragraphs[0].paragraph_format.space_before = Pt(1)

section(doc, "SELECTED AI ENGINEERING PROJECT")
p = doc.add_paragraph()
p.paragraph_format.keep_with_next = True
r = p.add_run("MediBot — Role-Aware Healthcare Knowledge Assistant")
r.bold = True
r.font.color.rgb = DARK
r = p.add_run(" | Personal Course Project")
r.italic = True
r.font.color.rgb = GRAY
bullet(doc, "Designed an advanced RAG solution that parses structured medical PDFs and Markdown using Docling and hierarchical, token-aware chunking while preserving headings and tables.")
bullet(doc, "Implemented server-side hybrid retrieval in Qdrant using dense vectors and BM25 sparse vectors with Reciprocal Rank Fusion, followed by cross-encoder reranking from the top 10 candidates to the top 3.")
bullet(doc, "Enforced role-based document access through Qdrant metadata filters at retrieval time for doctor, nurse, billing executive, technician, and administrator roles; validated against six adversarial prompts with no restricted collection leakage.")
bullet(doc, "Built SQL RAG over SQLite for analytical questions, including LLM-generated SQL sanitization, read-only execution, and natural-language result generation through a Groq cloud LLM API.")
bullet(doc, "Delivered a FastAPI backend and Streamlit interface with signed role tokens, source citations, retrieval-type indicators, and informative access-denied responses.")
p = doc.add_paragraph("Technologies: Python, FastAPI, Streamlit, Docling, Qdrant, FastEmbed, BM25, cross-encoder reranking, SQLite, Groq API")
p.paragraph_format.left_indent = Inches(0.2)
p.runs[0].bold = True
p.runs[0].font.size = Pt(9)
p.runs[0].font.color.rgb = GRAY

section(doc, "PROFESSIONAL EXPERIENCE")
role(doc, "Principal Software Engineer", "Dell Technologies", "June 2022 – Present")
bullet(doc, "Lead the end-to-end architecture and development of APOS and NPOS platforms, enabling reliable processing of approximately 120–150 order configurations per minute across multiple sales channels.")
bullet(doc, "Architect cloud-native microservices using Java, Spring Boot, REST APIs, and event-driven patterns to improve scalability, maintainability, and integration across the order ecosystem.")
bullet(doc, "Improved checkout performance by 40% and reduced latency by 35% by optimizing APIs, introducing effective caching strategies, and resolving application bottlenecks.")
bullet(doc, "Drive modernization of legacy APOS capabilities through service decomposition, technical-debt remediation, and incremental architectural improvements.")
bullet(doc, "Strengthen platform security and compliance through PCF deployment practices, OAuth-based access, engineering reviews, and production-readiness controls.")
bullet(doc, "Mentor engineers, lead design and code reviews, and collaborate with product, architecture, quality, and operations stakeholders to deliver resilient solutions.")

role(doc, "Senior Software Engineer", "Dell Technologies", "June 2019 – April 2022")
bullet(doc, "Led microservices development for the Financial Checks domain, delivering scalable and reliable validation services for critical business workflows.")
bullet(doc, "Reduced technical debt by 30% through refactoring, modernization, and maintainability improvements across core services.")
bullet(doc, "Improved financial-validation accuracy by 25% while enabling seamless integration with dependent enterprise systems.")
bullet(doc, "Supported SIT, UAT, production releases, and post-deployment stabilization to maintain high-quality, dependable delivery.")
bullet(doc, "Mentored team members and coordinated across engineering and business teams to resolve dependencies and improve execution efficiency.")

role(doc, "Technical Lead", "Collabera Technologies Pvt. Ltd. (Client: Dell Technologies)", "April 2017 – April 2019")
bullet(doc, "Owned end-to-end delivery of Incite, a product-configuration rules authoring application, with accountability for architecture, implementation, stability, and scalability.")
bullet(doc, "Led application modernization and technical-debt reduction initiatives that improved performance and long-term maintainability.")
bullet(doc, "Diagnosed and resolved complex cross-system issues across environments, strengthening reliability of integrated product-configuration workflows.")
bullet(doc, "Partnered with stakeholders and cross-functional teams to clarify requirements, manage dependencies, and improve delivery outcomes.")

role(doc, "Senior Software Developer & BI Consultant", "Creative Probers Software & Services Pvt. Ltd.", "February 2015 – July 2016")
bullet(doc, "Worked onsite at ITER in Saint-Paul-lès-Durance, France, to gather requirements, align technical solutions, and facilitate collaboration between teams in India and France.")
bullet(doc, "Developed SSIS packages and SSRS reports supporting data-warehouse integration, transformation, and business reporting requirements.")

role(doc, "Senior Software Developer", "Quantum BSO Tech Pvt. Ltd.", "December 2012 – January 2015")
bullet(doc, "Analyzed logistics-domain requirements and translated program specifications into maintainable user-interface and business-logic components.")
bullet(doc, "Contributed to application development, issue resolution, and delivery across the software development lifecycle.")

role(doc, "Senior Software Developer", "Syscon Infotech Pvt. Ltd.", "December 2007 – February 2009", "Mumbai, India")
bullet(doc, "Designed the application framework and developed user-interface and business-logic components for logistics solutions using ASP.NET and SQL Server.")

role(doc, "Software Developer", "Trust Systems & Software", "December 2006 – September 2007", "Nagpur, India")
bullet(doc, "Developed UI and application functionality using VB.NET and ASP.NET for banking invoicing and dormant-account maintenance solutions.")

section(doc, "EDUCATION & PROFESSIONAL DEVELOPMENT")
bullet(doc, "Bachelor of Engineering (Electrical Engineering), Nagpur University — 2003")
bullet(doc, "AI Engineering Program — In Progress; practical focus on RAG, embeddings, vector databases, hybrid search, reranking, SQL RAG, LLM integration, FastAPI, and Streamlit")
bullet(doc, "Ongoing application of GitHub Copilot and Windsurf for specification-driven development, refactoring, code review, and developer productivity")

# Header/footer
for section_obj in doc.sections:
    footer = section_obj.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Sushma Vaidya | Principal Software Engineer")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY

# Prevent isolated section headings where possible.
for p in doc.paragraphs:
    if p.style.name in {"Resume Section", "Resume Name"}:
        p.paragraph_format.keep_with_next = True

# Document metadata
doc.core_properties.title = "Sushma Vaidya - Principal Software Engineer Resume"
doc.core_properties.subject = ".NET, Java, Python, Cloud and AI Engineering"
doc.core_properties.author = "Sushma Vaidya"
doc.core_properties.keywords = ".NET, Java, Python, microservices, AI engineering, RAG, FastAPI, Qdrant, technical lead"

doc.save(OUT)
print(OUT)
