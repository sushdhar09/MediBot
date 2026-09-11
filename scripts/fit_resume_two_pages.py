from copy import deepcopy
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

SOURCE = r"C:\DellApplications\NewAssign Prep\Sushma_Dhar_Beautified_AI_DotNetLead_Resume.docx"
OUTPUT = r"C:\DellApplications\NewAssign Prep\Sushma_Dhar_Beautified_2Page_Resume.docx"


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    props.append(node)


def margins(cell, top=100, start=170, bottom=100, end=170):
    props = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    props.append(tc_mar)


def no_borders(table):
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    table._tbl.tblPr.append(borders)


def shell(doc):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.05)
    table.columns[1].width = Inches(5.4)
    no_borders(table)
    left, right = table.rows[0].cells
    left.width = Inches(2.05)
    right.width = Inches(5.4)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    shade(left, "17365D")
    shade(right, "FFFFFF")
    margins(left, 150, 190, 150, 190)
    margins(right, 70, 250, 70, 120)
    left._element.remove(left.paragraphs[0]._element)
    right._element.remove(right.paragraphs[0]._element)
    return left, right


def append_elements(cell, elements):
    for element in elements:
        cell._tc.append(deepcopy(element))


def paragraph_text(element):
    return "".join(element.itertext()) if element.tag == qn("w:p") else ""


doc = Document(SOURCE)
source_table = doc.tables[0]
left_elements = list(source_table.cell(0, 0)._tc)[1:-1]
right_elements = list(source_table.cell(0, 1)._tc)[1:-1]

right_split = next(
    i
    for i, element in enumerate(right_elements)
    if "Technical Lead" in paragraph_text(element) and "Collabera" in paragraph_text(element)
)
left_split = next(
    i
    for i, element in enumerate(left_elements)
    if "AI ENGINEERING" in paragraph_text(element)
)
page_one = right_elements[:right_split]
page_two = right_elements[right_split:]
left_page_one = left_elements[:left_split]
left_page_two = left_elements[left_split:]

body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

left_one, right_one = shell(doc)
append_elements(left_one, left_page_one)
append_elements(right_one, page_one)

page_break = doc.add_paragraph()
page_break.paragraph_format.space_after = 0
page_break.add_run().add_break(WD_BREAK.PAGE)

left_two, right_two = shell(doc)
append_elements(left_two, left_page_two)
append_elements(right_two, page_two)

doc.save(OUTPUT)
print(OUTPUT)
