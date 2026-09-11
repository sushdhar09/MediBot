from copy import deepcopy
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = r"C:\DellApplications\NewAssign Prep\Sushma_Dhar_Enhanced_AI_DotNetLead_Resume.docx"
OUTPUT = r"C:\DellApplications\NewAssign Prep\Sushma_Dhar_Beautified_AI_DotNetLead_Resume.docx"
NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
WHITE = RGBColor(255, 255, 255)
MUTED_WHITE = RGBColor(220, 230, 241)
DARK = RGBColor(37, 45, 55)
BLUE = RGBColor(31, 78, 121)
GRAY = RGBColor(89, 89, 89)


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    props.append(element)


def margins(cell, top=150, start=170, bottom=150, end=170):
    props = cell._tc.get_or_add_tcPr()
    tc_mar = props.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        props.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def remove_borders(table):
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    props.append(borders)


def sidebar_heading(cell, text):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(10.5)
    r.font.color.rgb = WHITE
    line = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "6EA5D7")
    line.append(bottom)
    p._p.get_or_add_pPr().append(line)


def sidebar_item(cell, title, detail=None):
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(title)
    r.bold = detail is not None
    r.font.name = "Aptos"
    r.font.size = Pt(9.2)
    r.font.color.rgb = WHITE
    if detail:
        r = p.add_run(f"\n{detail}")
        r.font.name = "Aptos"
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED_WHITE


def main_heading(cell, text):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(11.5)
    r.font.color.rgb = BLUE
    line = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "9EBCD5")
    line.append(bottom)
    p._p.get_or_add_pPr().append(line)


def main_bullet(cell, text):
    p = cell.add_paragraph(style="Resume Bullet")
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = DARK


def role(cell, title, company, dates, location="Bengaluru, India"):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{title} | {company}")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    r.font.color.rgb = DARK
    r = p.add_run(f"\n{dates} | {location}")
    r.italic = True
    r.font.name = "Aptos"
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY


doc = Document(SOURCE)
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

section = doc.sections[0]
section.top_margin = Inches(0.35)
section.bottom_margin = Inches(0.35)
section.left_margin = Inches(0.4)
section.right_margin = Inches(0.4)

shell = doc.add_table(rows=1, cols=2)
shell.alignment = WD_TABLE_ALIGNMENT.CENTER
shell.autofit = False
shell.columns[0].width = Inches(2.05)
shell.columns[1].width = Inches(5.4)
remove_borders(shell)
left, right = shell.rows[0].cells
left.width = Inches(2.05)
right.width = Inches(5.4)
left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
shade(left, NAVY)
shade(right, "FFFFFF")
margins(left, 190, 190, 190, 190)
margins(right, 100, 250, 100, 120)
left._element.remove(left.paragraphs[0]._element)
right._element.remove(right.paragraphs[0]._element)

p = left.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run("SV")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(27)
r.font.color.rgb = WHITE

p = left.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("PRINCIPAL SOFTWARE\nENGINEER")
r.bold = True
r.font.name = "Aptos"
r.font.size = Pt(9.5)
r.font.color.rgb = MUTED_WHITE

sidebar_heading(left, "Personal Information")
sidebar_item(left, "Location", "Bengaluru, India")
sidebar_item(left, "Phone", "+91 9886843824")
sidebar_item(left, "Email", "sush.dhar@gmail.com")

sidebar_heading(left, "Core Expertise")
for item in (
    ".NET & ASP.NET",
    "Java & Spring Boot",
    "Python",
    "Microservices",
    "REST APIs",
    "Event-Driven Systems",
    "Cloud Architecture",
    "SQL Server & MongoDB",
    "Kubernetes & CI/CD",
    "Performance Engineering",
):
    sidebar_item(left, f"• {item}")

sidebar_heading(left, "AI Engineering")
for item in (
    "Advanced RAG",
    "Dense + BM25 Search",
    "Qdrant Vector DB",
    "Cross-Encoder Reranking",
    "SQL RAG",
    "LLM & Groq APIs",
    "FastAPI & Streamlit",
    "Prompt Engineering",
):
    sidebar_item(left, f"• {item}")

sidebar_heading(left, "Leadership")
for item in (
    "Solution Architecture",
    "Technical Leadership",
    "Design & Code Reviews",
    "Engineering Mentorship",
    "Stakeholder Management",
    "Technical Debt Reduction",
):
    sidebar_item(left, f"• {item}")

sidebar_heading(left, "Education")
sidebar_item(left, "Bachelor of Engineering", "Electrical Engineering\nNagpur University, 2003")
sidebar_item(left, "AI Engineering Program", "In Progress")

p = right.add_paragraph()
p.paragraph_format.space_after = Pt(0)
r = p.add_run("SUSHMA VAIDYA")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(22)
r.font.color.rgb = BLUE
p = right.add_paragraph()
p.paragraph_format.space_after = Pt(5)
r = p.add_run("PRINCIPAL SOFTWARE ENGINEER  |  TECHNICAL LEAD")
r.bold = True
r.font.name = "Aptos"
r.font.size = Pt(10)
r.font.color.rgb = GRAY

main_heading(right, "PROFESSIONAL PROFILE")
p = right.add_paragraph(
    "Principal Software Engineer and technical leader with 14+ years of experience designing, modernizing, "
    "and delivering enterprise applications across .NET, Java, Python, microservices, and data platforms. "
    "Proven success leading mission-critical order-processing and financial-validation systems at Dell "
    "Technologies, improving checkout performance by 40%, reducing latency by 35%, and lowering technical "
    "debt by 30%. Experienced in cloud-native architecture, API design, security, CI/CD, performance "
    "engineering, stakeholder management, and mentoring. Currently expanding into AI engineering through "
    "hands-on work in advanced RAG, hybrid search, vector databases, reranking, SQL RAG, and LLM integration."
)
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.line_spacing = 1.03
for r in p.runs:
    r.font.size = Pt(10)
    r.font.color.rgb = DARK

main_heading(right, "SELECTED AI ENGINEERING PROJECT")
p = right.add_paragraph()
p.paragraph_format.keep_with_next = True
r = p.add_run("MediBot — Role-Aware Healthcare Knowledge Assistant")
r.bold = True
r.font.size = Pt(10.5)
r.font.color.rgb = DARK
r = p.add_run(" | Course Project")
r.italic = True
r.font.size = Pt(9)
r.font.color.rgb = GRAY
main_bullet(right, "Designed an advanced RAG solution using Docling and hierarchical, token-aware chunking to preserve document headings, tables, and structural context.")
main_bullet(right, "Implemented Qdrant hybrid retrieval with dense and BM25 sparse vectors, Reciprocal Rank Fusion, and cross-encoder reranking from top 10 to top 3.")
main_bullet(right, "Enforced vector-store RBAC for five staff roles and validated six adversarial prompts without restricted collection leakage.")
main_bullet(right, "Built read-only SQL RAG over SQLite plus a FastAPI backend and Streamlit UI with source citations and role-aware responses.")

main_heading(right, "PROFESSIONAL EXPERIENCE")
role(right, "Principal Software Engineer", "Dell Technologies", "June 2022 – Present")
main_bullet(right, "Lead end-to-end architecture and development of APOS and NPOS platforms, enabling reliable processing of approximately 120–150 order configurations per minute across sales channels.")
main_bullet(right, "Architect cloud-native microservices using Java, Spring Boot, REST APIs, and event-driven patterns for scalable enterprise integration.")
main_bullet(right, "Improved checkout performance by 40% and reduced latency by 35% through API optimization, caching, and bottleneck remediation.")
main_bullet(right, "Drive legacy APOS decomposition, technical-debt remediation, PCF/OAuth security practices, engineering reviews, and production readiness.")
main_bullet(right, "Mentor engineers and coordinate with product, architecture, quality, and operations stakeholders to deliver resilient solutions.")

role(right, "Senior Software Engineer", "Dell Technologies", "June 2019 – April 2022")
main_bullet(right, "Led microservices development for the Financial Checks domain, delivering scalable and reliable validation services for critical workflows.")
main_bullet(right, "Reduced technical debt by 30% and improved financial-validation accuracy by 25% through refactoring, modernization, and integration improvements.")
main_bullet(right, "Supported SIT, UAT, production releases, and post-deployment stabilization while mentoring engineers and coordinating cross-team dependencies.")

role(right, "Technical Lead", "Collabera Technologies Pvt. Ltd. (Client: Dell Technologies)", "April 2017 – April 2019")
main_bullet(right, "Owned end-to-end delivery of Incite, a product-configuration rules authoring application, covering architecture, implementation, stability, and scalability.")
main_bullet(right, "Led modernization and technical-debt reduction, resolved complex cross-system issues, and partnered with stakeholders to improve delivery outcomes.")

role(right, "Senior Software Developer & BI Consultant", "Creative Probers Software & Services Pvt. Ltd.", "February 2015 – July 2016")
main_bullet(right, "Worked onsite at ITER in Saint-Paul-lès-Durance, France, gathering requirements and coordinating delivery between teams in India and France.")
main_bullet(right, "Developed SSIS packages and SSRS reports supporting data-warehouse integration, transformation, and reporting.")

role(right, "Senior Software Developer", "Quantum BSO Tech Pvt. Ltd.", "December 2012 – January 2015")
main_bullet(right, "Translated logistics-domain requirements into maintainable user-interface and business-logic components across the delivery lifecycle.")

role(right, "Senior Software Developer", "Syscon Infotech Pvt. Ltd.", "December 2007 – February 2009", "Mumbai, India")
main_bullet(right, "Designed an application framework and developed UI and business-logic components using ASP.NET and SQL Server for logistics solutions.")

role(right, "Software Developer", "Trust Systems & Software", "December 2006 – September 2007", "Nagpur, India")
main_bullet(right, "Developed VB.NET and ASP.NET functionality for banking invoicing and dormant-account maintenance solutions.")

main_heading(right, "PROFESSIONAL DEVELOPMENT")
main_bullet(right, "AI Engineering Program — In Progress; hands-on focus on RAG, embeddings, vector databases, hybrid search, reranking, SQL RAG, LLM integration, FastAPI, and Streamlit.")
main_bullet(right, "Apply GitHub Copilot and Windsurf to specification-driven development, refactoring, code review, and developer productivity.")

for section_obj in doc.sections:
    footer = section_obj.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.clear()
    r = footer.add_run("Sushma Vaidya  |  Principal Software Engineer  |  sush.dhar@gmail.com")
    r.font.name = "Aptos"
    r.font.size = Pt(7.5)
    r.font.color.rgb = GRAY

doc.core_properties.title = "Sushma Vaidya - Principal Software Engineer Resume"
doc.core_properties.subject = ".NET, Java, Python, Cloud and AI Engineering"
doc.core_properties.author = "Sushma Vaidya"
doc.core_properties.keywords = ".NET, Java, Python, microservices, AI engineering, RAG, FastAPI, Qdrant, technical lead"
doc.save(OUTPUT)
print(OUTPUT)
